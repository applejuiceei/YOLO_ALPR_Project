from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\YOLO_ALPR_Project")
ASSET_ROOT = ROOT / "report_assets_sc850_progress" / "sc850_video71_vehicle_probe_20250626_102149"
OUTPUT = ROOT / "SC850SL_RK3588_项目进度汇报_20260711.docx"

FONT = "Microsoft YaHei"
MONO = "Consolas"
NAVY = "17365D"
BLUE = "2E74B5"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
LIGHT_GOLD = "FFF6DC"
GREEN = "2F6B45"
RED = "9B1C1C"


def set_font(run, name=FONT, size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def fixed_table(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for idx, width in enumerate(widths):
        if idx < len(grid.gridCol_lst):
            grid.gridCol_lst[idx].set(qn("w:w"), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, before=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 10)
    pf.space_after = Pt(7 if level == 1 else 5)
    set_keep_with_next(p)
    r = p.add_run(text)
    if level == 1:
        set_font(r, size=16, color=BLUE, bold=True)
    else:
        set_font(r, size=12.5, color=NAVY, bold=True)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, "F6F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_font(r, name=MONO, size=8.5, color="1F2937")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(label + "  ")
    set_font(r, size=10.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    fixed_table(table, [2700, 6660])
    for idx, (key, value) in enumerate(rows):
        cells = table.add_row().cells
        shade(cells[0], LIGHT_GRAY)
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(key), size=10, color=NAVY, bold=True)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    fixed_table(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        shade(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=9.5, color=NAVY, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(text), size=9.2)
            if idx in (0, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_image(doc, image_path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(7)
    r = cap.add_run(caption)
    set_font(r, size=9, color=MUTED, italic=True)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_font(r, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    r = paragraph.add_run(" 页")
    set_font(r, size=8.5, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("SC850SL / RK3588 实时道路车牌识别项目")
    set_font(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def main():
    doc = Document()
    configure_document(doc)

    # First page masthead
    add_para(doc, "技术项目进度汇报", size=10, color=BLUE, bold=True, after=4)
    title = add_para(doc, "SC850SL + RK3588\n实时道路车辆与车牌识别", size=23, color=NAVY, bold=True, after=5)
    title.paragraph_format.line_spacing = 1.02
    add_para(doc, "范围：远程 RK3588 板端 SC850SL 摄像头，基于 /dev/video71 的 NV12 实时道路画面", size=12, color=MUTED, after=12)
    add_kv_table(doc, [
        ("汇报日期", "2026-07-11"),
        ("当前目标", "在真实道路画面中完成车辆检测、真实车牌 OCR、投票锁定，并使车牌号随车辆显示。"),
        ("主控代码", "Python：rk3588_topk_capture_mipi_debug_sc850.py；Shell：run_sc850_video71_alpr_live.sh。"),
        ("当前状态", "核心闭环已在历史 1000 帧 SC850SL 路测中验证；最新常驻优化版因外部远程条件尚未完成板端回归。"),
    ])
    add_callout(
        doc,
        "汇报结论",
        "摄像头、ISP 输出、车辆检测、车牌 OBB、HyperLPR3 OCR、Top-K 与投票锁定已经在真实 /dev/video71 路测中形成闭环。当前工作重心已从“链路是否可用”转为“日间 ISP 曝光、近景 ROI、锁定速度、跟车稳定性与实时帧率”。",
        fill=LIGHT_BLUE,
    )
    add_image(doc, ASSET_ROOT / "frame_after.jpg", 6.55, "图 1  已完成 1000 帧真实 SC850SL 路测后的最终发布帧：可见车辆框、车牌框与锁定车牌文字。")

    add_heading(doc, "一、项目目标与实现路径", 1)
    add_para(doc, "最终目标不是“推流成功”或“显示 PLATE 占位”，而是在远程 RK3588 + SC850SL 真实道路相机上，让真实车牌号显示在对应车辆上方并随车移动。", size=10.5)
    add_matrix(doc, ["环节", "实现方式", "已验证情况"], [
        ("相机输入", "/dev/video71，3840x2160 NV12，V4L2 mmap", "已验证连续取帧"),
        ("车辆检测", "RKNN vehicle 模型，ROI 分块检测", "已验证真实道路车辆框"),
        ("车牌定位", "RKNN OBB 车牌模型", "已验证候选与几何过滤"),
        ("车牌识别", "HyperLPR3", "已验证真实 OCR 文本"),
        ("锁定与跟随", "Top-K、投票、Kalman/重关联", "历史路测已锁定；最新跟随策略待回归"),
        ("浏览器预览", "/tmp/frame.jpg + stream.py", "已验证浏览器显示"),
    ], [1700, 4600, 3060])

    add_heading(doc, "二、已完成工作与关键证据", 1)
    add_heading(doc, "1. 真实 SC850SL 摄像头链路恢复", 2)
    add_bullet(doc, "早期问题集中在 sensor/I2C/ISP 上游：I2C 未见 0x30、sensor id 为 000000、media graph 缺少 remote terminal sensor。")
    add_bullet(doc, "恢复后，驱动日志读取到 SC850SL id 009d1e；/dev/video71 可输出 4K NV12，RTSP 推流与 Python 直读均成功。")
    add_bullet(doc, "必须让 ALPR 独占 /dev/video71；mediamtx/GStreamer 推流会占用该设备，运行识别前需停止推流。")
    add_heading(doc, "2. 300 帧直读验证：先证明不是 OCR 问题", 2)
    add_matrix(doc, ["检查项", "结果", "结论"], [
        ("单帧 NV12", "12,441,600 bytes = 3840 x 2160 x 1.5", "Python 可直接读到完整原始帧"),
        ("连续车辆检测", "300 帧，约 1.32 FPS，1916 个车辆检测", "相机读取、NV12 解码、车辆 RKNN 可连续运行"),
        ("诊断模式", "OCR_ENGINE=none；min-process-vehicle-conf=1.1", "此阶段故意不进入 OCR，不能把 track_count=0 误解为识别失败"),
    ], [2050, 3650, 3660])
    add_heading(doc, "3. 1000 帧真实道路 ALPR 闭环", 2)
    add_matrix(doc, ["指标", "历史完整路测结果"], [
        ("处理帧数", "1000 帧"),
        ("耗时 / 平均速度", "1127.945 秒 / 约 0.89 FPS"),
        ("车辆检测", "10,237 原始候选；NMS 后 7,353"),
        ("有效车牌尝试", "2,831 次；OBB 候选 1,603 个"),
        ("HyperLPR3 OCR", "562 次"),
        ("投票锁定", "7 个锁定结果；无 OOM、无新的 sensor/ISP 错误"),
    ], [3000, 6360])
    add_para(doc, "代表性锁定结果包括：粤SEU850、苏EAZ8848、苏0CS180、苏E73X71、苏E73X7T、苏E8Z6J1、苏E2R91Q。", size=10.2, color=INK, after=5)

    images = doc.add_table(rows=1, cols=2)
    fixed_table(images, [4680, 4680])
    for cell in images.rows[0].cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = images.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSET_ROOT / "run" / "run_20250626_102153" / "track_729" / "full_rank1.jpg"), width=Inches(3.04))
    p = images.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSET_ROOT / "run" / "run_20250626_102153" / "track_729" / "plate_rank1.jpg"), width=Inches(2.50))
    captions = doc.add_table(rows=1, cols=2)
    fixed_table(captions, [4680, 4680])
    for idx, text in enumerate(["图 2  Track 729 的道路场景候选", "图 3  同一候选的车牌裁剪：苏E2R91Q"]):
        p = captions.cell(0, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=8.8, color=MUTED, italic=True)

    add_heading(doc, "三、遇到的困难、原因判断与处理", 1)
    add_matrix(doc, ["困难", "观察到的证据", "已采取措施 / 当前结论"], [
        ("SC850SL 上游 sensor 未上线", "I2C 无 0x30、sensor id 000000、remote terminal 失败", "重插/初始化恢复后，驱动识别到 id 009d1e，/dev/video71 已可用。"),
        ("推流与识别冲突", "mediamtx/GStreamer 和 Python 都需读取 /dev/video71", "演示/识别时停止推流；浏览器预览改由 /tmp/frame.jpg。"),
        ("4K Top-K 运行 OOM", "约 220 帧时 Python 被内核杀掉，匿名内存约 3.12 GiB", "将保留全帧缩到 960 宽、同帧候选共享预览、限制 rejected 样本；后续 1000 帧完整跑完。"),
        ("框太多、远车/电动车干扰", "低阈值与宽 ROI 会显示小目标和远处车辆", "新增近景 ROI、显示尺寸/面积/置信度门槛；最新版待板端回归。"),
        ("锁定后预测框漂移", "橙色/黄色框会在无真实检测时漂到空车道", "最新版关闭锁定预测显示，只保留短窗口真实检测重关联。"),
        ("日间 ISP 过曝", "白墙、车顶、斑马线高光溢出，车牌纹理减少", "先通过 v4l-subdev7 调低日间 exposure；候选 320/350/380，建议从 350 开始。"),
    ], [1900, 3150, 4310])
    add_callout(doc, "重要边界", "当前最新版本的显示过滤、右侧全 ROI、快速锁定、短窗口重关联和 1920 预览优化尚未在远程板上完成回归。它们是待验证改动，不应在汇报中表述为已达到的实验结果。", fill=LIGHT_GOLD, label_color="7A5A00")
    add_heading(doc, "四、最新版本说明（待板端回归）", 1)
    add_kv_table(doc, [
        ("主程序", "rk3588_topk_capture_mipi_debug_sc850.py：Python，负责 V4L2 取帧、RKNN、OBB、HyperLPR3、Top-K、投票、跟踪、预览输出。"),
        ("常驻启动脚本", "run_sc850_video71_alpr_live.sh：Shell，只负责固化运行参数并启动 Python 主程序。"),
        ("ROI", "0.03 0.32 1.00 0.95：覆盖右侧主车道，避开最远路口。"),
        ("锁定策略", "启用 rk_adaptive；同一真实车牌两帧一致即可锁定；重关联窗口缩短为 12 帧。"),
        ("显示策略", "仅显示满足近景尺寸、面积、置信度的车辆；关闭锁定车辆长时间预测框。"),
        ("性能策略", "1x2 ROI 检测；4K 原图仍用于识别；浏览器预览缩至 1920 宽、JPEG 质量 75。"),
    ])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "五、现场演示运行命令", 1)
    add_callout(doc, "演示前提", "远程板位于另一台电脑的局域网内，需要通过向日葵远程电脑操作。最新两个文件需先上传到 /root/alpr_topk_rk3588/。", fill=LIGHT_BLUE)
    add_heading(doc, "1. 停止旧识别与摄像头推流", 2)
    add_code(doc, "cd /root/alpr_topk_rk3588\nkill \"$(cat sc850_video71_alpr_live.pid)\" 2>/dev/null || true\ncd /root/mediamtx && ./stop_4_stream.sh")
    add_para(doc, "说明：如果识别是通过 nohup 启动的，Ctrl+C 只会停止 tail -f 日志跟随，不会停止后台识别进程。", size=9.8, color=RED, after=6)
    add_heading(doc, "2. 设置日间曝光（先读后设）", 2)
    add_code(doc, "v4l2-ctl -d /dev/v4l-subdev7 --get-ctrl=exposure,analogue_gain\nv4l2-ctl -d /dev/v4l-subdev7 --set-ctrl=analogue_gain=64\nv4l2-ctl -d /dev/v4l-subdev7 --set-ctrl=exposure=350")
    add_heading(doc, "3. 验证新文件与启动浏览器预览", 2)
    add_code(doc, "cd /root/alpr_topk_rk3588\nchmod +x run_sc850_video71_alpr_live.sh\npython3 -m py_compile rk3588_topk_capture_mipi_debug_sc850.py\nbash -n run_sc850_video71_alpr_live.sh\nsha256sum rk3588_topk_capture_mipi_debug_sc850.py run_sc850_video71_alpr_live.sh\n\n# 若 8080 没有预览服务，再另开一个终端执行\ncd /root/deploy\nnohup python3 stream.py > sc850_browser_stream.log 2>&1 &")
    add_heading(doc, "4. 启动常驻识别、查看日志、访问浏览器", 2)
    add_code(doc, "cd /root/alpr_topk_rk3588\nnohup ./run_sc850_video71_alpr_live.sh > sc850_video71_alpr_live.log 2>&1 &\necho $! > sc850_video71_alpr_live.pid\ntail -f sc850_video71_alpr_live.log\n\n# 浏览器： http://192.168.8.88:8080/")
    add_para(doc, "演示观察点：先看曝光是否保留车牌纹理，再看近景车辆框是否干净；随后观察 [LOCK]、[REASSOC] 日志，以及浏览器中车牌文字是否在同一车辆上持续出现。", size=10.2)

    next_heading = add_heading(doc, "六、下一步计划与验收标准", 1)
    next_heading.paragraph_format.page_break_before = True
    add_number(doc, "完成最新版板端回归：至少观察 300 帧实时浏览器效果，并在条件允许时完成 1000 帧完整采样包。")
    add_number(doc, "完成 ISP 日间参数小范围对比：exposure=320/350/380，固定 analogue_gain=64；以车牌纹理、高光溢出和锁定数量共同评估。")
    add_number(doc, "根据回归数据微调近景门槛：车辆尺寸、面积、置信度、ROI 上边界；目标是保留可读车牌的机动车，同时压低远车和电动车干扰。")
    add_number(doc, "验证锁定速度与跟车：记录“车辆首次进入 ROI -> 首次 OCR -> LOCK”的帧数，以及车牌文本在 track ID 变化后的重关联是否正确。")
    add_number(doc, "性能验收：记录实际 FPS、锁定延迟、CPU/NPU/内存；确认常驻运行不再出现 OOM，浏览器预览持续更新。")
    add_kv_table(doc, [
        ("成功标准", "近景机动车进入 ROI 后，在离开画面前完成真实车牌锁定；锁定文本显示在对应车辆上方，且不漂移到空车道或下一辆车。"),
        ("不接受的结果", "只显示 PLATE 占位、只有推流画面、OCR 未进入投票、浏览器显示旧帧、或把预测框当作真实跟踪。"),
        ("当前最大风险", "日间曝光与现场车牌清晰度；最新参数未回归；远程操作依赖外部电脑与网络环境。"),
    ])
    add_heading(doc, "七、汇报常见问答", 1)
    qa = [
        ("项目最终要交付什么？", "在远程 RK3588 板的 SC850SL 实时道路画面中，车辆框、真实车牌号和锁定状态能够同步显示，车牌号跟随对应车辆。"),
        ("目前是否已经跑通真实摄像头？", "是。历史 1000 帧 /dev/video71 路测已完成车辆检测、OBB、HyperLPR3、投票锁定和浏览器发布；但最新版优化参数尚未板端回归。"),
        ("为什么一开始推流不行，后来又可以？", "早期上游 sensor/I2C/ISP 未正常上线；恢复后驱动识别到 SC850SL id 009d1e，/dev/video71 能稳定输出 NV12。推流成功反映的是相机链路恢复，不等同于 OCR 已完成。"),
        ("为什么不用 C++ 重写？", "当前主线是 Python 调 RKNNLite、OpenCV 和 HyperLPR3。主要耗时在 4K 处理、NPU 推理和 OCR，不是 Python 语法本身；先调 ROI、ISP 和推理节奏收益更直接。"),
        ("0.89 FPS 是否太慢？", "是历史 4K 全量 2x2 分块、每帧检测和频繁 OCR 下的结果。最新版用 ROI、1x2、近景过滤和 1920 预览减负，实际加速倍数待板端测量。"),
        ("黄色框为什么会漂？", "它是锁定车辆在当前帧无真实检测时的 Kalman 预测框。最新版关闭长时间预测显示，并只在短窗口内做真实检测重关联。"),
        ("为什么 ISP 要先调？", "过曝会直接损失车牌纹理，后续 OCR 无法补回。日间应先选择合适 exposure，再讨论模型阈值和 OCR。"),
        ("最新版本现在能否直接宣称成功？", "不能。可以说明已完成本地语法检查和针对性优化；但必须等远程板端回归后，才能报告真实效果、FPS 和锁定延迟。"),
        ("下一步最关键的验证是什么？", "在日间合适曝光下，验证新 ROI 与快速锁定策略能否让近景车辆在离开画面前锁定，并确认文本不错误迁移到其他车辆。"),
    ]
    for question, answer in qa:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("问：" + question)
        set_font(r, size=10.5, color=NAVY, bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.15)
        r = p.add_run("答：" + answer)
        set_font(r, size=10.2)

    add_para(doc, "证据来源：远程板端采集包 sc850_video71_vehicle_probe_20250626_102149.tar.gz；该包内含 run_summary.json、候选图、frame_after.jpg、ALPR 日志及 dmesg。", size=8.5, color=MUTED, italic=True, after=0)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
